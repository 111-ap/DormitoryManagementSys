from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def read_sheet(
    content: bytes,
    required_headers: list[str],
    max_rows: int = 5000,
) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("无法读取 Excel 文件，请确认是 .xlsx 格式") from exc

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("Excel 文件为空") from exc

    headers = [str(x).strip() if x is not None else "" for x in raw_headers]
    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise ValueError(f"缺少表头：{'、'.join(missing)}")

    index = {h: headers.index(h) for h in required_headers}
    rows: list[dict[str, Any]] = []
    for row_no, values in enumerate(rows_iter, start=2):
        if len(rows) >= max_rows:
            raise ValueError(f"单次最多导入 {max_rows} 行")
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        item = {h: values[i] if i < len(values) else None for h, i in index.items()}
        item["_row"] = row_no
        rows.append(item)
    return rows, headers


def parse_gender(value: Any, default: str = "other") -> str:
    text = str(value or "").strip().lower()
    mapping = {
        "男": "male",
        "male": "male",
        "m": "male",
        "女": "female",
        "female": "female",
        "f": "female",
        "其他": "other",
        "other": "other",
    }
    return mapping.get(text, default)


def parse_bool(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return text in {"1", "true", "yes", "y", "是", "是舍长", "舍长"}


def parse_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"日期格式不正确：{text}")


def fmt_datetime(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else ""


def make_workbook(
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[int] | None = None,
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, start=1):
        width = widths[idx - 1] if widths and idx <= len(widths) else max(12, len(str(header)) + 4)
        ws.column_dimensions[get_column_letter(idx)].width = width
    return wb


def workbook_response(wb: Workbook, filename: str) -> StreamingResponse:
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
    }
    return StreamingResponse(stream, media_type=EXCEL_MEDIA_TYPE, headers=headers)
