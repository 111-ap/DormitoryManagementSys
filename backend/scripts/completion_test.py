"""附件、Excel 导入导出、学生账号创建的补充冒烟测试。"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app.core.config import settings  # noqa: E402
from app.main import app  # noqa: E402

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360000000020001e221bc330000000049454e44ae426082"
)


def must(resp, code=200, label=""):
    if resp.status_code != code:
        raise AssertionError(f"{label}: {resp.status_code} {resp.text}")
    return resp


def workbook_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> int:
    c = TestClient(app)
    login = must(
        c.post(
            "/api/v1/auth/login",
            json={"username": settings.admin_username, "password": settings.admin_initial_password},
        ),
        200,
        "admin login",
    ).json()
    h = {"Authorization": f"Bearer {login['access_token']}"}

    # 附件：先建报修单，再上传/下载/删除
    repair = must(
        c.post(
            "/api/v1/repairs",
            headers=h,
            json={"room_id": 1, "asset_id": 1, "description": "附件测试：灯管损坏"},
        ),
        200,
        "create repair",
    ).json()
    upload = must(
        c.post(
            "/api/v1/attachments/upload",
            headers=h,
            data={"biz_type": "repair", "biz_id": str(repair["id"])},
            files={"file": ("repair.png", PNG_1PX, "image/png")},
        ),
        200,
        "upload attachment",
    ).json()
    listing = must(
        c.get(f"/api/v1/attachments?biz_type=repair&biz_id={repair['id']}", headers=h),
        200,
        "list attachments",
    ).json()
    if len(listing) != 1:
        raise AssertionError("attachment list mismatch")
    download = must(c.get(upload["download_url"], headers=h), 200, "download attachment")
    if download.content != PNG_1PX:
        raise AssertionError("downloaded bytes mismatch")
    must(c.delete(f"/api/v1/attachments/{upload['id']}", headers=h), 200, "delete attachment")

    # Excel 学生导入
    students_xlsx = workbook_bytes(
        ["学号", "姓名", "性别", "院系编码", "班级", "手机号", "是否舍长"],
        [
            ["XL0001", "导入学生甲", "男", "CS", "软件 2601", "13800000001", "否"],
            ["XL0002", "导入学生乙", "男", "SE", "软件 2602", "13800000002", "是"],
        ],
    )
    imported = must(
        c.post(
            "/api/v1/excel/import/students",
            headers=h,
            files={"file": ("students.xlsx", students_xlsx, XLSX_MIME)},
        ),
        200,
        "import students",
    ).json()
    if imported["success"] != 2:
        raise AssertionError(f"student import mismatch: {imported}")

    # 导出学生
    exported = must(c.get("/api/v1/excel/export/students", headers=h), 200, "export students")
    wb = load_workbook(BytesIO(exported.content), read_only=True)
    ws = wb.active
    values = list(ws.iter_rows(values_only=True))
    if len(values) < 3:
        raise AssertionError("exported students rows missing")
    if values[0][0] != "学号":
        raise AssertionError("exported students header mismatch")

    # Excel 分配导入
    students = must(c.get("/api/v1/students?keyword=XL000", headers=h), 200, "list imported students").json()
    by_no = {x["student_no"]: x for x in students["items"]}
    assignments_xlsx = workbook_bytes(
        ["学号", "楼栋编码", "房间号", "床位号", "入住时间"],
        [
            ["XL0001", "B01", "101", "XL-B1", "2026-09-01 08:00:00"],
            ["XL0002", "B01", "101", "XL-B2", "2026-09-01 08:00:00"],
        ],
    )
    assigned = must(
        c.post(
            "/api/v1/excel/import/assignments",
            headers=h,
            files={"file": ("assignments.xlsx", assignments_xlsx, XLSX_MIME)},
        ),
        200,
        "import assignments",
    ).json()
    if assigned["success"] != 2:
        raise AssertionError(f"assignment import mismatch: {assigned}")

    # 学生账号创建 + 登录
    account = must(
        c.post(f"/api/v1/students/{by_no['XL0001']['id']}/account", headers=h, json={}),
        200,
        "create student account",
    ).json()
    if not account.get("initial_password"):
        raise AssertionError("generated initial password missing")
    stu_login = must(
        c.post(
            "/api/v1/auth/login",
            json={"username": account["username"], "password": account["initial_password"]},
        ),
        200,
        "student login",
    ).json()
    sh = {"Authorization": f"Bearer {stu_login['access_token']}"}
    me = must(c.get("/api/v1/students/me", headers=sh), 200, "student me").json()
    if me["student_no"] != "XL0001":
        raise AssertionError("student me mismatch")

    # 模板与其它导出
    must(c.get("/api/v1/excel/templates/students", headers=h), 200, "student template")
    must(c.get("/api/v1/excel/templates/assignments", headers=h), 200, "assignment template")
    must(c.get("/api/v1/excel/export/assignments", headers=h), 200, "export assignments")
    must(c.get("/api/v1/excel/export/repairs", headers=h), 200, "export repairs")
    must(c.get("/api/v1/excel/export/utility-bills", headers=h), 200, "export bills")
    must(c.get("/api/v1/excel/export/visitors", headers=h), 200, "export visitors")

    print("COMPLETION OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
